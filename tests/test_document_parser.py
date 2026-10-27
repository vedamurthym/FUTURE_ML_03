"""
Unit tests for resume document parsing module (src/document_parser.py).

Tests:
- Single PDF parsing (valid)
- Single DOCX parsing (valid)
- Multiple file parsing (1-20+ files)
- Unsupported file type handling (.txt, .png, etc.)
- Empty file handling (0 bytes)
- Corrupted file handling
- Mixture of valid and invalid files
- Text association with correct filename
- Maximum file size boundary check
"""

import io
import unittest
from pathlib import Path

import docx
import pypdf

from src.document_parser import (
    MAX_FILE_SIZE_BYTES,
    extract_text_from_docx,
    extract_text_from_pdf,
    format_file_size,
    parse_multiple_resumes,
    parse_single_resume,
)


def create_sample_docx(text: str = "Experienced Python developer with machine learning and SQL.") -> bytes:
    """Helper to generate in-memory DOCX bytes."""
    doc = docx.Document()
    doc.add_heading("Jane Developer", level=0)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_pdf(text: str = "Senior Engineer with Docker, Kubernetes, and Python expertise.") -> bytes:
    """Helper to generate in-memory valid PDF bytes."""
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    # We can write text or use a real PDF from archive
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


class TestDocumentParser(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.archive_pdf_path = Path("archive/data/data/ACCOUNTANT/10554236.pdf")

    def test_single_docx_parsing(self):
        """Test parsing a single valid DOCX document."""
        docx_bytes = create_sample_docx("Python, FastAPI, and PostgreSQL developer.")
        result = parse_single_resume("candidate.docx", docx_bytes)

        self.assertTrue(result["is_valid"])
        self.assertIsNone(result["error"])
        self.assertEqual(result["file_type"], "DOCX")
        self.assertEqual(result["filename"], "candidate.docx")
        self.assertIn("Python", result["extracted_skills"])
        self.assertIn("PostgreSQL", result["extracted_skills"])
        self.assertGreater(result["word_count"], 0)
        self.assertGreater(result["char_count"], 0)

    def test_single_pdf_parsing_from_archive(self):
        """Test parsing an actual PDF from the project archive."""
        if self.archive_pdf_path.exists():
            with open(self.archive_pdf_path, "rb") as f:
                pdf_bytes = f.read()

            result = parse_single_resume("10554236.pdf", pdf_bytes)
            self.assertTrue(result["is_valid"])
            self.assertEqual(result["file_type"], "PDF")
            self.assertEqual(result["filename"], "10554236.pdf")
            self.assertGreater(result["word_count"], 100)
            self.assertGreater(result["skill_count"], 0)

    def test_unsupported_file_type(self):
        """Test that non-PDF and non-DOCX files are rejected."""
        result = parse_single_resume("notes.txt", b"Plain text file")
        self.assertFalse(result["is_valid"])
        self.assertIn("Unsupported file type", result["error"])

        result2 = parse_single_resume("photo.png", b"\x89PNG\r\n\x1a\n")
        self.assertFalse(result2["is_valid"])
        self.assertIn("Unsupported file type", result2["error"])

    def test_empty_file_handling(self):
        """Test that zero-byte files return clear errors without crashing."""
        result_pdf = parse_single_resume("empty.pdf", b"")
        self.assertFalse(result_pdf["is_valid"])
        self.assertIn("empty", result_pdf["error"].lower())

        result_docx = parse_single_resume("empty.docx", b"")
        self.assertFalse(result_docx["is_valid"])
        self.assertIn("empty", result_docx["error"].lower())

    def test_corrupted_file_handling(self):
        """Test that corrupted bytes do not crash the parser."""
        corrupted_bytes = b"NOT_A_REAL_FILE_FORMAT_DATA_HERE"
        result_pdf = parse_single_resume("broken.pdf", corrupted_bytes)
        self.assertFalse(result_pdf["is_valid"])
        self.assertIn("corrupted", result_pdf["error"].lower())

        result_docx = parse_single_resume("broken.docx", corrupted_bytes)
        self.assertFalse(result_docx["is_valid"])
        self.assertIn("corrupted", result_docx["error"].lower())

    def test_max_file_size_exceeded(self):
        """Test that files exceeding 10 MB are rejected."""
        fake_large = b"x" * (MAX_FILE_SIZE_BYTES + 1024)
        result = parse_single_resume("large_resume.pdf", fake_large)
        self.assertFalse(result["is_valid"])
        self.assertIn("exceeds the maximum allowed limit", result["error"])

    def test_multiple_files_batch_and_correct_association(self):
        """Test parsing multiple files (PDFs and DOCX) with correct text association."""
        doc1 = create_sample_docx("Candidate One has skills in Python and SQL.")
        doc2 = create_sample_docx("Candidate Two has skills in Java and Docker.")
        doc3 = b"invalid content"

        file_items = [
            ("resume_1.docx", doc1),
            ("resume_2.docx", doc2),
            ("corrupted.pdf", doc3),
        ]

        batch = parse_multiple_resumes(file_items)
        self.assertEqual(batch["total_files"], 3)
        self.assertEqual(batch["valid_count"], 2)
        self.assertEqual(batch["error_count"], 1)

        # Check correct filename association
        f0 = batch["files"][0]
        self.assertEqual(f0["filename"], "resume_1.docx")
        self.assertTrue(f0["is_valid"])
        self.assertIn("Python", f0["extracted_skills"])

        f1 = batch["files"][1]
        self.assertEqual(f1["filename"], "resume_2.docx")
        self.assertTrue(f1["is_valid"])
        self.assertIn("Java", f1["extracted_skills"])

        f2 = batch["files"][2]
        self.assertEqual(f2["filename"], "corrupted.pdf")
        self.assertFalse(f2["is_valid"])

    def test_multiple_files_scaling_20_files(self):
        """Test batch parsing of 20 resumes."""
        file_items = []
        for i in range(20):
            doc = create_sample_docx(f"Candidate #{i} knows Python and Git.")
            file_items.append((f"candidate_{i}.docx", doc))

        batch = parse_multiple_resumes(file_items)
        self.assertEqual(batch["total_files"], 20)
        self.assertEqual(batch["valid_count"], 20)
        self.assertEqual(batch["error_count"], 0)


if __name__ == "__main__":
    unittest.main()
