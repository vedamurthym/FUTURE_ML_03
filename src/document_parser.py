"""
Resume Document Parser Module.

Safely parses uploaded PDF and DOCX resume documents in-memory.
Validates file extensions, size limits, and integrity.
Extracts raw text, normalizes text using src.text_preprocessor.clean_text,
and extracts skills using src.skill_extractor.extract_skills.
Never permanently saves uploaded files to disk or overwrites datasets.
"""

import io
import os
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
import pypdf
from pypdf.errors import PdfReadError

from src.text_preprocessor import clean_text
from src.skill_extractor import extract_skills

# Maximum file size per resume: 10 MB
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024
MAX_FILE_SIZE_MB = 10

ALLOWED_EXTENSIONS = {".pdf", ".docx"}


def format_file_size(size_bytes: int) -> str:
    """Format file size in bytes to human-readable string (KB or MB)."""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.2f} MB"


def extract_text_from_txt(content: bytes) -> str:
    """
    Extract readable text from plain text (TXT) bytes.
    Raises ValueError on empty or whitespace-only documents.
    """
    if not content:
        raise ValueError("Job description file is empty (0 bytes).")

    text = None
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        raise ValueError("Unable to decode text file. Ensure it is encoded in UTF-8 or standard text format.")

    text = text.strip()
    if not text:
        raise ValueError("Job description file contains no text (empty or whitespace only).")

    return text


def parse_job_description_file(filename: str, content: bytes) -> str:
    """
    Safely extract readable text from an uploaded job description file (TXT or PDF).
    Raises ValueError with a clear explanation if file is empty, corrupted, or unsupported.
    """
    clean_name = os.path.basename(filename or "job_description")
    ext = Path(clean_name).suffix.lower()

    if not content or len(content) == 0:
        raise ValueError(f"Job description file '{clean_name}' is empty (0 bytes).")

    if ext == ".txt":
        return extract_text_from_txt(content)
    elif ext == ".pdf":
        return extract_text_from_pdf(content)
    else:
        raise ValueError(
            f"Unsupported job description format '{ext or 'unknown'}'. Only TXT (.txt) and PDF (.pdf) files are supported."
        )


def extract_text_from_pdf(content: bytes) -> str:
    """
    Extract readable text from PDF bytes.
    Raises ValueError on corrupted, empty, or encrypted documents.
    """
    if not content:
        raise ValueError("PDF content is empty.")

    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
    except (PdfReadError, Exception) as e:
        raise ValueError("Corrupted, malformed, or unreadable PDF file.")

    if reader.is_encrypted:
        raise ValueError("PDF is encrypted or password-protected.")

    if len(reader.pages) == 0:
        raise ValueError("PDF has no pages.")

    page_texts = []
    for idx, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
            if txt.strip():
                page_texts.append(txt.strip())
        except Exception:
            # Continue extracting text from other pages if one page has issues
            continue

    combined_text = "\n\n".join(page_texts).strip()
    if not combined_text:
        raise ValueError("No extractable text found in PDF (may contain only scanned images or blank pages).")

    return combined_text


def extract_text_from_docx(content: bytes) -> str:
    """
    Extract readable text from DOCX bytes.
    Raises ValueError on corrupted or invalid documents.
    """
    if not content:
        raise ValueError("DOCX content is empty.")

    try:
        doc = docx.Document(io.BytesIO(content))
    except (zipfile.BadZipFile, Exception) as e:
        raise ValueError("Corrupted, malformed, or unreadable DOCX file.")

    parts = []
    # 1. Paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)

    # 2. Tables (resumes often use tables for contact info or experience)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))

    combined_text = "\n".join(parts).strip()
    if not combined_text:
        raise ValueError("No extractable text found in DOCX (document contains no text content).")

    return combined_text


def parse_single_resume(filename: str, content: bytes) -> Dict[str, Any]:
    """
    Safely validate and parse a single resume file.
    Does not crash on invalid/corrupted files; returns structured status dictionary.
    """
    clean_filename = os.path.basename(filename or "unnamed_resume")
    file_ext = Path(clean_filename).suffix.lower()
    file_size = len(content) if content else 0
    formatted_size = format_file_size(file_size)

    # Base failure result template
    def make_error_result(msg: str, doc_type: str = "UNKNOWN") -> Dict[str, Any]:
        return {
            "filename": clean_filename,
            "is_valid": False,
            "error": msg,
            "file_size": file_size,
            "file_size_formatted": formatted_size,
            "file_type": doc_type,
            "raw_text": "",
            "text_preview": "",
            "cleaned_text": "",
            "word_count": 0,
            "char_count": 0,
            "extracted_skills": [],
            "skills_by_category": {},
            "skill_count": 0,
        }

    # 1. Check extension
    if not file_ext or file_ext not in ALLOWED_EXTENSIONS:
        return make_error_result(
            f"Unsupported file type '{file_ext or 'none'}'. Only PDF and DOCX files are supported."
        )

    doc_type = "PDF" if file_ext == ".pdf" else "DOCX"

    # 2. Check empty file
    if file_size == 0:
        return make_error_result("File is empty (0 bytes).", doc_type)

    # 3. Check file size limit
    if file_size > MAX_FILE_SIZE_BYTES:
        return make_error_result(
            f"File size ({formatted_size}) exceeds the maximum allowed limit of {MAX_FILE_SIZE_MB} MB.",
            doc_type,
        )

    # 4. Extract text according to format
    try:
        if file_ext == ".pdf":
            raw_text = extract_text_from_pdf(content)
        else:
            raw_text = extract_text_from_docx(content)
    except ValueError as ve:
        return make_error_result(str(ve), doc_type)
    except Exception:
        return make_error_result(f"Failed to process {doc_type} file due to unreadable content.", doc_type)

    # 5. Clean text & extract skills using existing core NLP modules
    cleaned = clean_text(raw_text)
    skill_result = extract_skills(raw_text)
    extracted_skills = skill_result.get("skills", [])
    skills_by_category = skill_result.get("by_category", {})

    preview = raw_text[:400].strip()
    if len(raw_text) > 400:
        preview += "..."

    return {
        "filename": clean_filename,
        "is_valid": True,
        "error": None,
        "file_size": file_size,
        "file_size_formatted": formatted_size,
        "file_type": doc_type,
        "raw_text": raw_text,
        "text_preview": preview,
        "cleaned_text": cleaned,
        "word_count": len(raw_text.split()),
        "char_count": len(raw_text),
        "extracted_skills": extracted_skills,
        "skills_by_category": skills_by_category,
        "skill_count": len(extracted_skills),
    }


def parse_multiple_resumes(file_items: List[Tuple[str, bytes]]) -> Dict[str, Any]:
    """
    Parse a collection of (filename, bytes) items.
    Returns aggregated summary and per-file results.
    """
    results = []
    valid_count = 0
    error_count = 0

    for filename, content in file_items:
        res = parse_single_resume(filename, content)
        results.append(res)
        if res["is_valid"]:
            valid_count += 1
        else:
            error_count += 1

    return {
        "total_files": len(file_items),
        "valid_count": valid_count,
        "error_count": error_count,
        "files": results,
    }
